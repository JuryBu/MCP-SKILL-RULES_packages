import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.dont_write_bytecode = True
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / 'src' / 'inspector'))

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

import pptx_inspector as inspector


class PptxInspectionV2(unittest.TestCase):
    def setUp(self):
        self.temporary = tempfile.TemporaryDirectory(prefix='pptx-inspection-v2-')
        self.path = Path(self.temporary.name) / 'neutral.pptx'
        self.presentation = Presentation()
        self.slide = self.presentation.slides.add_slide(self.presentation.slide_layouts[6])

    def tearDown(self):
        self.temporary.cleanup()

    def text(self, left=1, top=1, width=3, height=1, text='Neutral words', size=18, shapes=None):
        shape = (shapes if shapes is not None else self.slide.shapes).add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height))
        shape.text = text
        frame = shape.text_frame
        frame.auto_size = MSO_AUTO_SIZE.NONE
        frame.word_wrap = True
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(size)
        return shape

    def rectangle(self, left=.8, top=.8, width=3.5, height=1.5, opacity=1, geometry=MSO_SHAPE.RECTANGLE):
        shape = self.slide.shapes.add_shape(geometry, Inches(left), Inches(top), Inches(width), Inches(height))
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor(255, 255, 255)
        shape.line.fill.background()
        if opacity != 1:
            color = shape._element.xpath('./p:spPr/a:solidFill/*')[0]
            alpha = OxmlElement('a:alpha')
            alpha.set('val', str(round(opacity * 100000)))
            color.append(alpha)
        return shape

    def detect(self, checks=None):
        self.presentation.save(self.path)
        return inspector.detect_issues(str(self.path), checks=checks or ['overlap', 'overflow', 'readability', 'alignment'])

    def test_opaque_foreground_detected(self):
        self.text()
        self.rectangle()
        result = self.detect(['overlap'])
        self.assertEqual(result['summary']['errors'], 1)
        self.assertEqual(result['issues'][0]['metadata']['assessment'], 'confirmed')
        self.assertTrue(result['structure'][0]['metadata']['inspectionLimitations'])

    def test_background_preserved_without_alarm(self):
        self.rectangle()
        self.text()
        result = self.detect(['overlap'])
        self.assertEqual(result['issues'], [])
        self.assertEqual([element['type'] for element in result['structure'][0]['elements']], ['shape', 'text'])

    def test_transparent_decoration_not_mask(self):
        self.text()
        self.rectangle(opacity=0)
        self.rectangle(opacity=.35)
        self.assertEqual(self.detect(['overlap'])['issues'], [])

    def test_nonrectangular_drawing_not_confirmed_mask(self):
        self.text()
        self.rectangle(geometry=MSO_SHAPE.ARC)
        self.assertEqual(self.detect(['overlap'])['summary']['errors'], 0)

    def test_filled_oval_over_text_keeps_nonempty_candidate(self):
        self.text(left=2, top=1.3, width=1, height=.5, text='Visible text')
        self.rectangle(geometry=MSO_SHAPE.OVAL)
        issues = self.detect(['overlap'])['issues']
        self.assertTrue(issues)
        self.assertTrue(all(issue['metadata']['assessment'] == 'candidate' for issue in issues))

    def test_oval_behind_text_is_normal_composition(self):
        self.rectangle(geometry=MSO_SHAPE.OVAL)
        self.text(left=2, top=1.3, width=1, height=.5, text='Visible text')
        self.assertEqual(self.detect(['overlap'])['issues'], [])

    def test_no_fill_stroke_not_cover(self):
        self.text()
        self.rectangle().fill.background()
        self.assertEqual(self.detect(['overlap'])['issues'], [])

    def test_wrapped_text_fits(self):
        self.text(width=2, height=2, text='Neutral words for a normal wrapped paragraph', size=18)
        self.assertEqual(self.detect(['readability'])['issues'], [])

    def test_cjk_normal_wrap_fits(self):
        self.text(width=2, height=2, text='中性示例文字用于检查正常自动换行布局', size=18)
        self.assertEqual(self.detect(['readability'])['issues'], [])

    def test_clipping_is_candidate(self):
        self.text(width=1, height=.15, text='Neutral words exceed a short box', size=20)
        issues = self.detect(['readability'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['type'], 'clipped')
        self.assertEqual(issues[0]['metadata']['assessment'], 'candidate')
        self.assertNotEqual(issues[0]['severity'], 'error')

    def test_no_wrap_width_candidate(self):
        self.text(width=1, height=1, text='A long unwrapped neutral sentence').text_frame.word_wrap = False
        self.assertIn('right', self.detect(['readability'])['issues'][0]['metadata']['sides'])

    def test_autofit_not_clipping(self):
        for mode in (MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE, MSO_AUTO_SIZE.SHAPE_TO_FIT_TEXT):
            with self.subTest(mode=mode):
                shape = self.text(width=1, height=.1, text='Many neutral words with automatic sizing')
                shape.text_frame.auto_size = mode
        self.assertEqual(self.detect(['readability'])['issues'], [])

    def test_text_frame_whitespace_only_low_candidate(self):
        self.text(top=1, height=2, text='Top')
        self.text(top=2, height=1, text='Bottom')
        issues = self.detect(['overlap'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['severity'], 'info')
        self.assertEqual(issues[0]['metadata']['confidence'], 'low')
        self.assertIn('estimated_lines_separate', issues[0]['metadata']['reasonCodes'])

    def test_actual_text_overlap_remains_candidate(self):
        self.text(text='First line')
        self.text(top=1.05, text='Second line')
        issues = self.detect(['overlap'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['metadata']['reasonCodes'], ['estimated_text_line_intersection'])
        self.assertEqual(issues[0]['metadata']['assessment'], 'candidate')

    def test_margins_and_center_anchor(self):
        shape = self.text(height=2, text='Centered')
        shape.text_frame.margin_left = Inches(.3)
        shape.text_frame._txBody.bodyPr.set('anchor', 'ctr')
        element = self.detect(['overlap'])['structure'][0]['elements'][0]
        region = element['metadata']['textRegions'][0]
        self.assertAlmostEqual(region['x0'], Inches(1.3))
        self.assertGreater(region['y0'], Inches(1.5))

    def test_decorative_bleed_vs_text_overflow(self):
        self.rectangle(left=-1)
        self.text(left=-1, text='Outside page')
        issues = self.detect(['overflow'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['elements'][0]['type'], 'text')

    def test_nested_group_coordinates(self):
        outer = self.slide.shapes.add_group_shape()
        inner = outer.shapes.add_group_shape()
        self.text(left=1, top=2, width=2, height=1, text='Nested', shapes=inner.shapes)
        for group, offset, extent, child_offset, child_extent in (
            (inner, (3, 4), (4, 2), (1, 2), (2, 1)),
            (outer, (1, 1), (4, 2), (3, 4), (4, 2)),
        ):
            transform = group._element.xpath('./p:grpSpPr/a:xfrm')[0]
            for node_name, values, attributes in (('off', offset, ('x', 'y')), ('ext', extent, ('cx', 'cy')),
                                                  ('chOff', child_offset, ('x', 'y')), ('chExt', child_extent, ('cx', 'cy'))):
                node = transform.find('{http://schemas.openxmlformats.org/drawingml/2006/main}' + node_name)
                for attribute, value in zip(attributes, values):
                    node.set(attribute, str(Inches(value)))
        result = self.detect(['overlap'])
        elements = result['structure'][0]['elements']
        self.assertEqual([element['type'] for element in elements], ['group', 'group', 'text'])
        self.assertEqual(elements[-1]['bounds'], {'x0': Inches(1), 'y0': Inches(1), 'x1': Inches(5), 'y1': Inches(3)})
        self.assertEqual(result['issues'], [])
        found = inspector.search_text_region(str(self.path), 'Nested')
        self.assertTrue(found['found'])
        self.assertEqual(found['rect'], elements[-1]['bounds'])

    def test_group_rotation_maps_child_bounds(self):
        group = self.slide.shapes.add_group_shape()
        self.text(left=1, top=1, width=2, height=1, text='Rotated', shapes=group.shapes)
        group.rotation = 90
        element = self.detect(['overlap'])['structure'][0]['elements'][-1]
        for key, value in {'x0': 1.5, 'y0': .5, 'x1': 2.5, 'y1': 2.5}.items():
            self.assertAlmostEqual(element['bounds'][key], Inches(value), places=5)

    def test_group_child_remains_occludable(self):
        group = self.slide.shapes.add_group_shape()
        self.text(shapes=group.shapes)
        self.rectangle()
        self.assertEqual(self.detect(['overlap'])['summary']['errors'], 1)

    def test_unresolved_font_is_candidate_not_error(self):
        first = self.text()
        second = self.text()
        first.text_frame.paragraphs[0].font.size = None
        second.text_frame.paragraphs[0].font.size = None
        issues = self.detect(['overlap'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['metadata']['confidence'], 'low')
        self.assertEqual(issues[0]['metadata']['assessment'], 'candidate')

    def test_cover_with_unresolved_font_is_candidate(self):
        self.text().text_frame.paragraphs[0].font.size = None
        self.rectangle()
        issues = self.detect(['overlap'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['metadata']['assessment'], 'candidate')

    def test_cover_with_unsupported_vertical_text_is_candidate(self):
        self.text().text_frame._txBody.bodyPr.set('vert', 'vert')
        self.rectangle()
        issues = self.detect(['overlap'])['issues']
        self.assertEqual(len(issues), 1)
        self.assertEqual(issues[0]['metadata']['assessment'], 'candidate')

    def test_rotated_rectangle_not_strong_cover(self):
        self.text()
        self.rectangle().rotation = 45
        result = self.detect(['overlap'])
        self.assertEqual(result['summary']['errors'], 0)
        self.assertTrue(result['issues'])
        self.assertTrue(all(issue['metadata']['assessment'] == 'candidate' for issue in result['issues']))

    def test_local_spacing_positive(self):
        for left in (1, 2.3, 4.5):
            self.rectangle(left=left, top=3, width=1, height=.5)
        issues = self.detect(['alignment'])['issues']
        self.assertTrue(any(issue['type'] == 'uneven-spacing' for issue in issues))

    def test_unrelated_columns_not_grouped(self):
        for left, size, width in ((.2, 10, .3), (1, 10, .3), (4, 18, 1.5), (6.5, 18, 1.5)):
            self.text(left=left, top=3, width=width, height=.5, text='Label', size=size)
        self.assertEqual(self.detect(['alignment'])['issues'], [])

    def test_tiled_cells_allow_different_column_widths(self):
        for left, width in ((1, 1), (2, 1.4), (3.4, 1)):
            self.rectangle(left=left, top=3, width=width, height=.5)
        self.assertEqual(self.detect(['alignment'])['issues'], [])

    def test_independent_card_size_anomaly_retained(self):
        for left, width in ((1, 1), (2.3, 1.4), (4, 1)):
            self.rectangle(left=left, top=3, width=width, height=.5)
        self.assertTrue(any(issue['type'] == 'inconsistent-size' for issue in self.detect(['alignment'])['issues']))

    def test_cli_json_compatibility(self):
        self.text()
        self.presentation.save(self.path)
        request = {'action': 'detect_issues', 'pptxPath': str(self.path), 'page': 1}
        process = subprocess.run([sys.executable, str(Path(inspector.__file__))],
            input=json.dumps(request), text=True, encoding='utf-8', capture_output=True,
            env={**os.environ, 'PYTHONIOENCODING': 'utf-8', 'PYTHONDONTWRITEBYTECODE': '1'}, timeout=20)
        self.assertEqual(process.returncode, 0, process.stdout + process.stderr)
        result = json.loads(process.stdout)
        self.assertTrue(result['ok'])
        self.assertEqual(set(result['result']), {'summary', 'issues', 'structure'})

    def test_zero_alpha_text_not_confirmed_cover(self):
        shape = self.text()
        shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 0, 0)
        color = shape._element.xpath('.//a:defRPr/a:solidFill/*')[0]
        alpha = OxmlElement('a:alpha')
        alpha.set('val', '0')
        color.append(alpha)
        self.rectangle()
        self.assertEqual(self.detect(['overlap'])['issues'], [])

    def test_issue_limit_is_explicit_and_not_empty(self):
        for index in range(5):
            self.text(text=f'Neutral {index}')
        with patch('inspection_geometry.MAX_INSPECTION_ISSUES', 3):
            result = self.detect(['overlap'])
        self.assertEqual(len(result['issues']), 3)
        budget = result['structure'][0]['metadata']['inspectionBudget']
        self.assertTrue(budget['issueLimitReached'])
        self.assertTrue(budget['notExhaustive'])

    def test_pair_limit_explains_zero_issues(self):
        for index in range(6):
            self.rectangle(left=index, top=3, width=.5, height=.5)
        with patch('inspection_geometry.MAX_COMPARISON_PAIRS', 4):
            result = self.detect(['overlap'])
        self.assertEqual(result['issues'], [])
        budget = result['structure'][0]['metadata']['inspectionBudget']
        self.assertEqual(budget['pairChecksPerformed'], 4)
        self.assertTrue(budget['pairLimitReached'])
        self.assertTrue(budget['notExhaustive'])

    def test_text_line_cap_reports_incomplete_evidence(self):
        self.text(text='First\nSecond\nThird', height=3)
        with patch('pptx_evidence.MAX_TEXT_LINES', 2):
            result = self.detect(['overlap'])
        self.assertTrue(result['structure'][0]['elements'][0]['metadata']['textLayoutLimited'])
        self.assertTrue(result['structure'][0]['metadata']['inspectionBudget']['detailLimitReached'])


if __name__ == '__main__':
    unittest.main()
